
import os
import re
import shutil
from datetime import datetime
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.hyperlink import Hyperlink

import insumos
from consolidar_tablas import VERSION_TABLAS, tabla_valor_vigente
from tabla_construccion import (COMUNAS_7, COMUNAS_10, COMUNAS_5, COMUNAS_5N,
                                COMUNAS_FALTANTES)

RAIZ = Path(os.environ.get("PRELIQ_RAIZ") or Path(__file__).resolve().parent.parent)
CARPETA_HOMOLOGACION = RAIZ / "input" / "tablas"
HOJA = "USOS"
# Tono clarito por categoria en las columnas del uso (A a H).
COLOR_CATEGORIA = {"residencial": "E8F1FB", "comercial": "FDF0E4",
                   "industrial": "F0EAF8", "institucional": "E8F5EC",
                   "anexos": "F2F2F2"}
# El mismo uso escrito distinto en la base y en el libro de homologacion.
ALIAS_USOS = {"Saunas_y_Banos_Turcos": "Saunas_y_Banios_Turcos"}
HOJA_CONV = "CONVENCIONALES"
CARPETA_SALIDA = RAIZ / "results" / "HOMOLOGACION"

# Tablas que se reparten por grupo de comunas y tipologia de la ZHF.
TIPOLOGIAS = {
    "T1_RESIDENCIAL": ["011", "012", "013", "014", "015", "016"],
    "T2_EDIFICIOS": ["011", "012", "013", "014", "015", "016"],
    "T3_COMERCIAL": ["021", "022", "023"],
    "T4_INDUSTRIAL": ["031", "032", "033"],
}
# Desde la version 2, EDIFICIOS e INDUSTRIAL parten las 10 comunas en 5C y 5C_N.
POR_5C = {"T2_EDIFICIOS", "T4_INDUSTRIAL"}
# Tabla unica para las 17 comunas.
TABLAS_17C = ["T5_INSTITUCIONAL_ED", "T6_INSTITUCIONAL_SA",
              "T8_INSTITUCIONAL_IG", "T9_HOTELES",
              "T11_CCOMERCIALES", "T13_UNIDAD_DEPORTIVA"]
# Pensiones y residencias (038): tabla fija, ver Liquidacion_tablas.py.
TABLA_FIJA = {"Pensiones_y_Residencias": ["023"]}

COMUNAS_17 = sorted(set(COMUNAS_7) | set(COMUNAS_10) - set(COMUNAS_FALTANTES))
EXTRA = sorted(COMUNAS_FALTANTES)
ORDEN_GRUPO = {"7C": 0, "10C": 1, "5C": 2, "5C_N": 3, "17C": 4, "": 5}

NOTAS = {
    "MODELO": "Se liquida por modelo; con el modelo apagado queda en $ 0.",
    "ESPECIALES": "Valor especial: no sale de tabla.",
    "SIN TABLA": "Sin tabla de valor asignada.",
    "T10_ANEXOS": "Anexo: la T10 no está aprobada, no se revaloriza.",
    "T12_PARQUEADEROS": "Parqueaderos: regla propia (T12), fuera del consolidado.",
}


def comunas_de(grupo: str) -> str:
    """Las comunas del grupo, sin las 5 extra: por ahora no se trabajan."""
    base = {"7C": COMUNAS_7, "5C": COMUNAS_5, "17C": COMUNAS_17,
            "10C": [c for c in COMUNAS_10 if c not in EXTRA],
            "5C_N": [c for c in COMUNAS_5N if c not in EXTRA]}.get(grupo, [])
    return ", ".join(sorted(base))


def grupo_de(base: str, comuna: str) -> str:
    if comuna in COMUNAS_7:
        return "7C"
    if base in POR_5C:
        return "5C" if comuna in COMUNAS_5 else "5C_N"
    return "10C"


def familia(tabla: str) -> str:
    """T1_RESIDENCIAL_013_9 -> T1_RESIDENCIAL; T5_INSTITUCIONAL_ED se queda."""
    for base in list(TIPOLOGIAS) + TABLAS_17C:
        if tabla == base or tabla.startswith(base + "_"):
            return base
    m = re.match(r"^(T\d+_[A-Z]+(?:_[A-Z]+)?)", tabla)
    return m.group(1) if m and tabla.startswith(("T7", "T8")) else tabla


def descripcion_usos() -> pd.DataFrame:
    """Lo descriptivo de la hoja Uso del libro de homologacion, por USO_LADM."""
    libros = sorted(CARPETA_HOMOLOGACION.glob(
        "HOMOLOGACION_Y_TABLAS_DE_VALOR_USOS_LADM_*.xlsx"))
    libros = [f for f in libros if not f.name.startswith("~$")]
    if not libros:
        print("   (sin libro de homologacion en input/tablas: la hoja USOS sale "
              "sin la parte descriptiva)")
        return pd.DataFrame()
    u = pd.read_excel(libros[-1], sheet_name="Uso", header=1, dtype=str)
    u = u.rename(columns={u.columns[0]: "CATEGORÍA", u.columns[1]: "CÓDIGO ANTERIOR",
                          u.columns[2]: "NOMBRE DESTINO",
                          u.columns[3]: "CÓDIGO HOMOLOGACIÓN",
                          u.columns[4]: "USO_LADM", u.columns[5]: "TIPO",
                          u.columns[6]: "DESCRIPCIÓN"})
    # Las celdas combinadas llegan vacias debajo de la primera.
    u[["CATEGORÍA", "USO_LADM", "CÓDIGO HOMOLOGACIÓN"]] = u[
        ["CATEGORÍA", "USO_LADM", "CÓDIGO HOMOLOGACIÓN"]].ffill()
    u = u[u["USO_LADM"].notna()]
    u["USO_LADM"] = u["USO_LADM"].str.strip()
    u["CÓDIGO HOMOLOGACIÓN"] = u["CÓDIGO HOMOLOGACIÓN"].str.strip()

    def unir(s):
        vals = [str(x).strip() for x in s.dropna() if str(x).strip()
                and str(x).strip().upper() != "N/A"]
        return " / ".join(dict.fromkeys(vals))
    # Un mismo nombre de uso puede tener dos codigos (Coliseos: 056 y 114):
    # cada par uso-codigo es un uso distinto.
    return (u.groupby(["USO_LADM", "CÓDIGO HOMOLOGACIÓN"], sort=False)
             .agg({"CATEGORÍA": "first", "CÓDIGO ANTERIOR": unir,
                   "NOMBRE DESTINO": unir,
                   "TIPO": lambda s: unir(s), "DESCRIPCIÓN": "first"})
             .reset_index())


def motivo_no_liquidado(uso: str, cod: str, unico: bool = True) -> str:
    """Por que un uso del libro no aparece en la liquidacion: no esta en la
    base, solo esta en zona rural, o esta pero va a preinconsistencias."""
    if "base" not in _CACHE_ENERO:
        ruta = Path(insumos.base_construcciones(str(RAIZ / "output" / "todas_las_construcciones.parquet")))
        if ruta.exists():
            b = pd.read_parquet(ruta, columns=["USO_LADM", "DESTINOCONS",
                                               "COMUNA"])
            b["USO_LADM"] = (b["USO_LADM"].astype(str).str.strip()
                             .replace(ALIAS_USOS))
            b["DEST"] = b["DESTINOCONS"].astype(str).str.zfill(3)
            b["URBANA"] = (b["COMUNA"].astype(str).str.strip().str.zfill(2)
                           .isin(COMUNAS_17 + EXTRA))
            _CACHE_ENERO["base"] = b
        else:
            _CACHE_ENERO["base"] = None
    b = _CACHE_ENERO["base"]
    if b is None:
        return "No aparece en la liquidación"
    x = b[(b["USO_LADM"] == uso) & (b["DEST"] == cod)]
    if x.empty and unico:            # anexos: el destino viene en 000
        x = b[b["USO_LADM"] == uso]
    if x.empty:
        return "No aparece en la liquidación: no existe en la base"
    urb = int(x["URBANA"].sum())
    if urb == 0:
        return (f"No aparece en la liquidación: solo en zona rural "
                f"({len(x):,} construcciones)")
    return (f"No aparece en la liquidación: está en la base ({urb:,} "
            f"construcciones urbanas) pero va a preinconsistencias")


def completar_usos(t: pd.DataFrame, mapa: dict) -> pd.DataFrame:
    """Agrega los usos del libro de homologacion que no aparecen en la
    liquidacion, armados desde su REGLA: la tabla y la condicion de enero, los
    grupos de comunas de hoy y el enlace al bloque. Es un manual: cada uso
    lleva su regla aunque en la base no tenga construcciones."""
    desc = descripcion_usos()
    if desc.empty:
        return t
    presentes = set(zip(t["USO_LADM"], t["_COD"]))
    usos_presentes = set(t["USO_LADM"])

    def llave(uso, fam, g, cond9):
        # Solo las tablas por grupo distinguen grupo y, la T1, condicion 9.
        con_grupo = fam in TIPOLOGIAS or fam in TABLAS_17C or fam == "T12_PARQUEADEROS"
        return (uso, fam, g if con_grupo else "*",
                bool(cond9) if fam == "T1_RESIDENCIAL" else False)
    opciones = {llave(u, str(tb).split(" · ")[0], g, c9)
                for u, tb, g, c9 in zip(t["USO_LADM"], t["TABLA DE VALOR"],
                                        t["GRUPO"], t["_COND9"])}
    codigos_por_uso = desc.groupby("USO_LADM")["CÓDIGO HOMOLOGACIÓN"].nunique()
    reglas = {}
    for r in _REGLAS_ENERO():
        reglas.setdefault(r["uso"], []).append(r)
    nuevas, vistos = [], set()
    for uso, cod in zip(desc["USO_LADM"], desc["CÓDIGO HOMOLOGACIÓN"]):
        ya_esta = (uso, cod) in presentes or (uso in usos_presentes
                                              and codigos_por_uso[uso] == 1)
        if ya_esta and uso not in reglas:
            continue
        filas_uso = reglas.get(uso) or [{"fam": "SIN TABLA",
                                         "cond": set(TODAS_COND),
                                         "com_enero": set()}]
        for r in filas_uso:
            fam = r["fam"]
            cond9 = r["cond"] == {9}
            grupos = _grupos_nuevos(fam, {c.lstrip("0") for c in r["com_enero"]})
            for g in grupos or [""]:
                clave = (uso, cod, fam, g, cond9)
                if clave in vistos or (ya_esta and llave(uso, fam, g, cond9)
                                       in opciones):
                    continue
                vistos.add(clave)
                enlaces = []
                tabla = fam
                if fam in TIPOLOGIAS and g:
                    for tip in TABLA_FIJA.get(uso, TIPOLOGIAS[fam]):
                        origen = f"{fam}_{tip}" + ("_9" if cond9 else "")
                        enlaces.append((tip, mapa.get((origen, g))))
                    if fam == "T1_RESIDENCIAL" and g == "10C":
                        tabla += " · COND_9" if cond9 else " · COND_0"
                elif fam in TABLAS_17C:
                    enlaces.append(("17C", mapa.get((fam, "17C"))))
                nuevas.append({"USO_LADM": uso,
                               "CONDICIÓN": condicion_aplicacion_enero(
                                   uso, fam, g, cond9),
                               "TABLA DE VALOR": tabla, "GRUPO": g,
                               "COMUNAS": comunas_de(g) if g else "",
                               "CONSTRUCCIONES": 0, "OBSERVACIÓN": "",
                               "_ENLACES": enlaces,
                               "_ORDEN": ORDEN_GRUPO.get(g, 9),
                               "_COND9": cond9, "_MOD": False, "_COD": cod})
    if not nuevas:
        return t
    print(f"   Hoja {HOJA}: se agregan {len(nuevas)} filas de regla sin "
          f"construcciones en la liquidación.")
    return pd.concat([t, pd.DataFrame(nuevas)], ignore_index=True)


def condicion_general(uso: str) -> str:
    """La CONDICION (posicion 22 del numero predial) del uso, tal cual la
    escribe la columna H de la hoja Uso de enero."""
    if "npn" not in _CACHE_ENERO:
        libros = [f for f in sorted(CARPETA_HOMOLOGACION.glob(
            "HOMOLOGACION_Y_TABLAS_DE_VALOR_USOS_LADM_*.xlsx"))
            if not f.name.startswith("~$")]
        u = pd.read_excel(libros[-1], sheet_name="Uso", header=1,
                          dtype=str).iloc[:, [4, 7]]
        u.columns = ["USO", "COND"]
        u["USO"] = u["USO"].ffill().str.strip()
        u = u.dropna(subset=["COND"])
        _CACHE_ENERO["npn"] = (u.groupby("USO")["COND"].first()
                               .map(lambda x: " ".join(str(x).split())))
    return _CACHE_ENERO["npn"].get(uso, "")


def condicion_aplicacion_enero(uso, fam, grupo, cond9) -> str:
    """La CONDICION APLICACION TABLA de enero para esta fila: la regla del
    mismo uso y tabla en el grupo de comunas que le corresponde (7C = las
    '3 y 22' de enero; 10C, 5C y 5C_N = las '1,9,10,11 y 12')."""
    viejas = ({"03", "22"} if grupo == "7C" else
              {"01", "09", "10", "11", "12"} if grupo in ("10C", "5C", "5C_N")
              else None)
    if (uso, grupo) in CORRECCIONES_CONDICION:
        return CORRECCIONES_CONDICION[(uso, grupo)]
    # Usos de modelo: la condicion 9 va por MODELO y la 8 por su tabla.
    try:
        from comparacion_vigencia import CONFIG as CONFIG_VIG
        por_modelo = CONFIG_VIG.get("usos_por_modelo", {})
    except Exception:                                    # pragma: no cover
        por_modelo = {}
    if uso in por_modelo and fam != "MODELO":
        return str(por_modelo[uso])
    rs = [r for r in _REGLAS_ENERO() if r["uso"] == uso and r["fam"] == fam]
    if not rs:
        return ""
    sel = [r for r in rs if viejas is None or (r["com_enero"] & viejas)] or rs
    if fam == "T1_RESIDENCIAL" and grupo == "10C" and len(sel) > 1:
        por_cond = [r for r in sel if r["cond"] != "RESTO" and
                    ((9 in r["cond"]) == cond9) and
                    (not cond9 or r["cond"] == {9})]
        sel = por_cond or [r for r in sel if r["cond"] == "RESTO"] or sel
    # El texto de enero tal cual; vacio si enero lo deja vacio.
    return "; ".join(dict.fromkeys(r["texto_condap"] for r in sel
                                   if r["texto_condap"]))


_CACHE_ENERO = {}

# Correcciones a la CONDICION APLICACION TABLA de enero. En las 7 comunas la
# T1 no distingue por condicion: enero escribia "Condicion del 1 al 9" y
# "Condiciones 1 a 9", que dejaban por fuera la condicion 0.
CORRECCIONES_CONDICION = {("Barracas", "7C"): "Todas",
                          ("Jardin_Infantil_en_Casa", "7C"): "Todas"}


def _REGLAS_ENERO():
    if "r" not in _CACHE_ENERO:
        _CACHE_ENERO["r"] = reglas_enero()
    return _CACHE_ENERO["r"]


def condicion_regla(uso, base, grupo, cond9, mod, por_modelo, reglas_txt) -> str:
    """La condicion juridica que dice la regla para esta fila. Lo que el
    codigo decide por condicion va primero; lo demas, como lo escribe la hoja
    Uso del libro de homologacion; sin regla escrita, 'Todas'."""
    if base == "MODELO" or mod:
        return "9"
    try:
        from pre_inconsistencias import USOS_PAR_DE_MODELO
    except Exception:                                    # pragma: no cover
        USOS_PAR_DE_MODELO = set()
    if uso in USOS_PAR_DE_MODELO:
        return "Diferente de 9"
    if uso in por_modelo:
        return str(por_modelo[uso])                      # la que va por tabla
    if base == "T1_RESIDENCIAL" and grupo == "10C":
        return "9" if cond9 else "Diferente de 9"
    if base == "T1_RESIDENCIAL":                         # 7C: una sola hoja
        return "Todas"
    textos = [t for t in reglas_txt.get((uso, base), [])
              if t and t.lower() not in ("nan", "none")]
    if textos:
        return "; ".join(dict.fromkeys(" ".join(t.split()) for t in textos))
    return "Todas"


def filas_usos(liq: pd.DataFrame, mapa: dict) -> pd.DataFrame:
    """Una fila por uso, tabla, grupo de comunas y condicion (9 o no 9)."""
    cols = [c for c in ("USO_LADM", "CONDICION", "COMUNA", "TABLA_ORIGEN",
                        "DESTINOCONS", "DESTANEX") if c in liq.columns]
    d = liq[cols].copy()
    d["COMUNA"] = d["COMUNA"].astype(str).str.strip().str.zfill(2)
    d = d[d["COMUNA"].isin(COMUNAS_17 + EXTRA)]           # sin rurales
    dest = (d["DESTINOCONS"].astype(str).str.strip().str.zfill(3)
            if "DESTINOCONS" in d else pd.Series("", index=d.index))
    anex = (d["DESTANEX"].astype(str).str.strip().str.zfill(3)
            if "DESTANEX" in d else pd.Series("", index=d.index))
    d["_COD"] = dest.where(~dest.isin(["000", "nan", "<NA>", "None"]), anex)
    d["TABLA_ORIGEN"] = d["TABLA_ORIGEN"].astype(str)
    d["USO_LADM"] = d["USO_LADM"].astype(str).str.strip().replace(ALIAS_USOS)
    d["CONDICION"] = d["CONDICION"].astype(str)
    d["BASE"] = d["TABLA_ORIGEN"].map(familia)
    d["COND9"] = d["TABLA_ORIGEN"].str.endswith("_9")
    # Garajes, parqueaderos y depositos (T12) se asignan solo por destino: van
    # en las 17 comunas, como las tablas de bloque unico.
    d["GRUPO"] = [grupo_de(b, c) if b in TIPOLOGIAS
                  else ("17C" if b in TABLAS_17C or b == "T12_PARQUEADEROS"
                        else "")
                  for b, c in zip(d["BASE"], d["COMUNA"])]

    try:
        from comparacion_vigencia import CONFIG as CONFIG_VIG
        por_modelo = CONFIG_VIG.get("usos_por_modelo", {})
    except Exception:                                    # pragma: no cover
        por_modelo = {}
    # Lo que va por MODELO (salvo la condicion que va por tabla). Con el modelo
    # apagado la liquidacion le pone la tabla de su familia, y la comparacion
    # de vigencias lo excluye: se marca aparte para que la hoja diga lo que es.
    d["MOD"] = [u in por_modelo and c != str(por_modelo[u])
                for u, c in zip(d["USO_LADM"], d["CONDICION"])]

    # La condicion que dice la REGLA (no la observada en los datos).
    reglas_txt = {}
    for r in reglas_enero():
        reglas_txt.setdefault((r["uso"], r["fam"]), []).append(r["texto_cond"])

    filas = []
    for (uso, base, grupo, cond9, mod), s in d.groupby(
            ["USO_LADM", "BASE", "GRUPO", "COND9", "MOD"], sort=False):
        condicion = condicion_aplicacion_enero(uso, base, grupo, cond9)

        enlaces = []                                     # (texto, columna)
        if base in TIPOLOGIAS:
            for tip in TABLA_FIJA.get(uso, TIPOLOGIAS[base]):
                origen = f"{base}_{tip}" + ("_9" if cond9 else "")
                col = mapa.get((origen, grupo))
                enlaces.append((tip, col))
            # En las 10 comunas la T1 tiene dos hojas: COND_9 y COND_0.
            tabla = base + ((" · COND_9" if cond9 else " · COND_0")
                            if base == "T1_RESIDENCIAL" and grupo == "10C"
                            else "")
        elif base in TABLAS_17C:
            enlaces.append(("17C", mapa.get((base, "17C"))))
            tabla = base
        else:
            tabla = base

        nota = NOTAS.get(base, "")
        if base.startswith(("T7", "T8")) and not any(c for _, c in enlaces):
            nota = "Esta tabla no viene en el consolidado de esta versión."
        if base in TIPOLOGIAS or base in TABLAS_17C:
            if not any(c for _, c in enlaces):
                nota = "No se encontró la columna en el consolidado."
        if uso in por_modelo:
            extra = (f"En la comparación de vigencias se trata como modelo "
                     f"salvo condición {por_modelo[uso]}.")
            nota = f"{nota} {extra}".strip()
        if uso in TABLA_FIJA:
            nota = f"Tabla fija {base}_{TABLA_FIJA[uso][0]} en todas las comunas."
        if base == "MODELO":
            tabla, enlaces = "MODELO", []                # sin hipervinculo
        elif mod:
            # Uso de modelo pero fuera de las comunas del modelo: va por tabla
            # en la liquidacion, y la comparacion de vigencias lo excluye.
            tabla = f"{tabla} · fuera de las comunas del modelo (en la app se excluye)"

        filas.append({"USO_LADM": uso, "CONDICIÓN": condicion,
                      "TABLA DE VALOR": tabla, "GRUPO": grupo,
                      "COMUNAS": comunas_de(grupo),
                      "CONSTRUCCIONES": len(s), "OBSERVACIÓN": nota,
                      "_ENLACES": enlaces, "_ORDEN": ORDEN_GRUPO.get(grupo, 9),
                      "_COND9": cond9, "_MOD": mod,
                      "_COD": s["_COD"].mode().iat[0] if len(s) else ""})
    return pd.DataFrame(filas)


def agregar_hoja_usos(liq: pd.DataFrame | None = None,
                      ruta_consolidado: str | None = None,
                      ruta_salida: str | None = None) -> str | None:
    """Arma HOMOLOGACION_Y_TABLAS_DE_VALOR_<version>_<fecha>.xlsx: una copia
    del consolidado vigente con la hoja USOS al frente. El consolidado no se
    toca. Devuelve la ruta del libro nuevo."""
    origen = Path(ruta_consolidado) if ruta_consolidado else tabla_valor_vigente()
    if origen is None or not Path(origen).exists():
        print("   (no hay consolidado de tablas: la hoja USOS no se arma)")
        return None
    fecha = datetime.now().strftime("%Y%m%d")
    ruta = (Path(ruta_salida) if ruta_salida else CARPETA_SALIDA /
            f"HOMOLOGACION_Y_TABLAS_DE_VALOR_{VERSION_TABLAS}_{fecha}.xlsx")
    ruta.parent.mkdir(parents=True, exist_ok=True)
    try:
        shutil.copyfile(origen, ruta)
    except PermissionError:
        raise RuntimeError(f"No se pudo escribir '{ruta}'. Ciérrelo en Excel y "
                           f"vuelva a ejecutar.")
    if liq is None:
        from comparacion_vigencia import CONFIG as CONFIG_VIG
        liq = pd.read_parquet(CONFIG_VIG["parquet_liquidacion"],
                              columns=["USO_LADM", "CONDICION", "COMUNA",
                                       "TABLA_ORIGEN", "DESTINOCONS",
                                       "DESTANEX"])
    from comparacion_vigencia import _mapa_tablas_valor
    mapa = _mapa_tablas_valor(str(ruta))

    t = filas_usos(liq, mapa)
    revisar_reglas(t, mapa)
    t = completar_usos(t, mapa)
    dif_enero = comparar_con_enero(liq, mapa)
    if not dif_enero.empty:
        print(f"\n   REVISAR LAS REGLAS: {len(dif_enero)} diferencia(s) con la hoja "
              f"Uso de enero:")
        for f in dif_enero.head(15).itertuples(index=False):
            print(f"     - {f.USO_LADM}: {f.DIFERENCIA.lower()} (enero "
                  f"{f[3] or '-'}, liquidación {f[4]}; {f.CONSTRUCCIONES:,} "
                  f"construcciones)")
    desc = descripcion_usos()
    if not desc.empty:
        por_par = desc.rename(columns={"CÓDIGO HOMOLOGACIÓN": "_COD"})
        t = t.merge(por_par, on=["USO_LADM", "_COD"], how="left")
        # Si el codigo observado no esta en el libro, se toma el del uso.
        por_uso = desc.drop_duplicates("USO_LADM").set_index("USO_LADM")
        falta = t["CATEGORÍA"].isna()
        for c in por_uso.columns:
            destino = c if c != "CÓDIGO HOMOLOGACIÓN" else "_COD_LIBRO"
            t.loc[falta, destino] = t.loc[falta, "USO_LADM"].map(por_uso[c])
        t["CÓDIGO HOMOLOGACIÓN"] = t.get("_COD_LIBRO", pd.Series(index=t.index)
                                         ).fillna(t["_COD"].where(~falta))
    for c in ["CATEGORÍA", "CÓDIGO ANTERIOR", "NOMBRE DESTINO",
              "CÓDIGO HOMOLOGACIÓN", "TIPO", "DESCRIPCIÓN"]:
        if c not in t.columns:
            t[c] = ""
    # "ANEXOS" viene en mayusculas en el libro de homologacion: "Anexos".
    t["CATEGORÍA"] = t["CATEGORÍA"].map(
        lambda v: v.capitalize() if isinstance(v, str) and v.isupper() else v)
    # TIPO por el codigo de homologacion: hasta el 076 son convencionales.
    codigo = pd.to_numeric(t["CÓDIGO HOMOLOGACIÓN"], errors="coerce")
    t["TIPO"] = [("CONVENCIONAL" if n <= 76 else "NO CONVENCIONAL")
                 if pd.notna(n) else
                 ("NO CONVENCIONAL" if cat == "Anexos" else "")
                 for n, cat in zip(codigo, t["CATEGORÍA"])]
    t["_CODIGO"] = t["CÓDIGO HOMOLOGACIÓN"].fillna("999")
    t = t.sort_values(["_CODIGO", "USO_LADM", "TABLA DE VALOR", "_ORDEN",
                       "_COND9"], ascending=[True, True, True, True, False])

    t["CONDICIÓN GENERAL"] = t["USO_LADM"].map(condicion_general)
    columnas = ["CATEGORÍA", "CÓDIGO ANTERIOR", "NOMBRE DESTINO",
                "CÓDIGO HOMOLOGACIÓN", "USO_LADM", "TIPO", "DESCRIPCIÓN",
                "CONDICIÓN GENERAL", "CONDICIÓN", "TABLA DE VALOR", "GRUPO",
                "COMUNAS", "HIPERVÍNCULO"]
    # Encabezados como en la hoja Uso del libro de enero.
    titulos = {"CONDICIÓN GENERAL": "CONDICIÓN (POSICIÓN 22 NÚMERO PREDIAL "
                                    "NACIONAL)",
               "TABLA DE VALOR": "TABLA DE VALOR O MODELO",
               "CONDICIÓN": "CONDICIÓN APLICACIÓN TABLA"}
    # Las columnas del uso (A a H) se combinan sobre sus filas; el color del
    # bloque va de la tabla en adelante.
    n_uso = columnas.index("CONDICIÓN GENERAL") + 1

    wb = load_workbook(ruta)
    if HOJA in wb.sheetnames:
        del wb[HOJA]
    ws = wb.create_sheet(HOJA, 0)
    bloques = encabezar_bloques(wb[HOJA_CONV])
    bloque_t12 = hoja_t12(wb)

    azul = PatternFill("solid", fgColor="83CAEB")
    borde = Border(*(Side(style="thin", color="B7C9D6"),) * 4)
    negrilla = Font(bold=True, color="0F1F33")

    ws.cell(1, 1, f"USOS LADM Y TABLAS DE VALOR · versión {VERSION_TABLAS} · "
                  f"{Path(ruta).name}").font = Font(bold=True, size=13,
                                                     color="0F1F33")
    ws.merge_cells(start_row=1, start_column=1, end_row=1,
                   end_column=len(columnas))
    for j, nombre in enumerate(columnas, start=1):
        c = ws.cell(2, j, titulos.get(nombre, nombre))
        c.fill, c.font, c.border = azul, negrilla, borde
        c.alignment = Alignment(horizontal="center", vertical="center",
                                wrap_text=True)

    n_links = 0
    for i, fila in enumerate(t.itertuples(index=False), start=3):
        datos = dict(zip(t.columns, fila))
        # El bloque del consolidado al que va la fila: el de su primera columna.
        cols = [c for _, c in datos["_ENLACES"] if c]
        bloque = bloques.get(prefijo_bloque(cols[0])) if cols else None
        if (bloque is None and bloque_t12
                and datos["TABLA DE VALOR"] == "T12_PARQUEADEROS"):
            bloque = bloque_t12
        fondo = PatternFill("solid", fgColor=bloque["claro"]) if bloque else None
        tono = COLOR_CATEGORIA.get(str(datos.get("CATEGORÍA") or "").strip().lower())
        fondo_cat = PatternFill("solid", fgColor=tono) if tono else None
        for j, nombre in enumerate(columnas, start=1):
            v = datos.get(nombre)
            c = ws.cell(i, j, None if (v is None or (not isinstance(v, list)
                                                     and pd.isna(v))) else v)
            c.border = borde
            c.alignment = Alignment(vertical="center", wrap_text=True)
            if fondo is not None and j > n_uso:
                c.fill = fondo
            elif fondo_cat is not None and j <= n_uso:
                c.fill = fondo_cat
        if bloque:
            c = ws.cell(i, len(columnas), f"{bloque['etiqueta']} →")
            c.hyperlink = Hyperlink(ref=c.coordinate,
                                    location=f"'{bloque.get('hoja', HOJA_CONV)}'!"
                                             f"{bloque['celda']}")
            c.fill = PatternFill("solid", fgColor=bloque["color"])
            c.font = Font(bold=True, underline="single",
                          color=bloque["texto"])
            c.alignment = Alignment(horizontal="center", vertical="top")
            n_links += 1

    # Celdas combinadas por uso (codigo + nombre), como en el libro de enero.
    llave = list(zip(t["CÓDIGO HOMOLOGACIÓN"].astype(str), t["USO_LADM"]))
    fila0 = 3
    for k in range(1, len(llave) + 1):
        if k == len(llave) or llave[k] != llave[fila0 - 3]:
            fin = k + 2
            if fin > fila0:
                for j in range(1, n_uso + 1):
                    ws.merge_cells(start_row=fila0, start_column=j,
                                   end_row=fin, end_column=j)
            fila0 = fin + 1

    anchos = [14, 13, 30, 13, 34, 18, 46, 30, 30, 22, 8, 34, 30]
    for j, a in enumerate(anchos, start=1):
        ws.column_dimensions[get_column_letter(j)].width = a
    ws.row_dimensions[2].height = 32
    ws.freeze_panes = "F3"
    ws.auto_filter.ref = f"A2:{get_column_letter(len(columnas))}{len(t) + 2}"
    n_filas = len(t)

    # El manual no lleva CONTROL (se queda en el consolidado) ni la hoja de
    # diferencias: esas se avisan en pantalla al correr el proceso.
    for hoja in ("CONTROL", HOJA_REVISAR):
        if hoja in wb.sheetnames:
            del wb[hoja]
    try:
        wb.save(ruta)
    except PermissionError:
        raise RuntimeError(f"No se pudo guardar '{ruta}'. Ciérrelo en Excel y "
                           f"vuelva a ejecutar.")
    print(f"   Homologación y tablas de valor: {ruta} (hoja {HOJA}: {n_filas} "
          f"filas, {n_links} hipervínculos)")
    return str(ruta)


ETIQUETAS = {"T1_RESIDENCIAL": "RESIDENCIAL", "T2_EDIFICIOS": "EDIFICIOS",
             "T3_COMERCIAL": "COMERCIAL", "T4_INDUSTRIAL": "INDUSTRIAL",
             "T5_INSTITUCIONAL_ED": "INSTITUCIONAL EDUCATIVO",
             "T6_INSTITUCIONAL_SA": "INSTITUCIONAL SALUD",
             "T8_INSTITUCIONAL_IG": "INSTITUCIONAL IGLESIAS",
             "T9_HOTELES": "HOTELES", "T11_CCOMERCIALES": "CENTROS COMERCIALES",
             "T13_UNIDAD_DEPORTIVA": "UNIDAD DEPORTIVA"}


def prefijo_bloque(columna: str) -> str:
    """T2_EDIFICIOS_5C_N_013 -> T2_EDIFICIOS_5C_N; T9_HOTELES_17C se queda."""
    return re.sub(r"_\d{3}$", "", str(columna))


def encabezar_bloques(conv) -> dict:
    """Pone en CONVENCIONALES una fila arriba que combina las columnas de cada
    tabla (RESIDENCIAL · 7C, EDIFICIOS · 5C_N...) con su color. Devuelve, por
    bloque, la celda combinada y sus colores, para enlazarla desde USOS."""
    from consolidar_tablas import FAMILIAS, _mezclar, _texto_contraste

    encabezados = [(c.column, str(c.value)) for c in conv[1] if c.value]
    conv.insert_rows(1)
    conv.cell(1, 1, "PUNTAJE")
    conv.merge_cells(start_row=1, start_column=1, end_row=2, end_column=1)
    a1 = conv.cell(1, 1)
    a1.fill = PatternFill("solid", fgColor="2C3E50")
    a1.font = Font(bold=True, color="FFFFFF")
    a1.alignment = Alignment(horizontal="center", vertical="center")

    # Columnas seguidas con el mismo prefijo forman un bloque.
    bloques, actual = {}, None
    for col, nombre in encabezados[1:]:                  # la 1a es PUNTAJE
        pref = prefijo_bloque(nombre)
        if actual and actual["prefijo"] == pref and col == actual["fin"] + 1:
            actual["fin"] = col
            continue
        actual = {"prefijo": pref, "inicio": col, "fin": col}
        bloques[pref] = actual

    for pref, b in bloques.items():
        fam = next((f for f in sorted(ETIQUETAS, key=len, reverse=True)
                    if pref.startswith(f)), pref)
        grupo = pref[len(fam):].strip("_").replace("_COND_", " · COND_")
        b["etiqueta"] = f"{ETIQUETAS.get(fam, fam)} · {grupo}" if grupo else \
            ETIQUETAS.get(fam, fam)
        base = next((FAMILIAS[p] for p in sorted(FAMILIAS, key=len, reverse=True)
                     if pref.startswith(p)), "5D6D7E")
        b["color"] = _mezclar(base, 0.15)
        b["claro"] = _mezclar(base, 0.85)
        b["texto"] = _texto_contraste(b["color"])
        b["celda"] = f"{get_column_letter(b['inicio'])}1"
        c = conv.cell(1, b["inicio"], b["etiqueta"])
        if b["fin"] > b["inicio"]:
            conv.merge_cells(start_row=1, start_column=b["inicio"], end_row=1,
                             end_column=b["fin"])
        c.fill = PatternFill("solid", fgColor=b["color"])
        c.font = Font(bold=True, color=b["texto"])
        c.alignment = Alignment(horizontal="center", vertical="center")
    conv.row_dimensions[1].height = 24
    conv.freeze_panes = "B3"

    # Y de vuelta: un enlace a USOS al final de la fila de bloques.
    volver = conv.cell(1, max(b["fin"] for b in bloques.values()) + 2,
                       "← Volver a USOS")
    volver.hyperlink = Hyperlink(ref=volver.coordinate, location=f"'{HOJA}'!A1")
    volver.font = Font(bold=True, color="0563C1", underline="single")
    return bloques


LUGARES_REGLAS = [
    "src/app_vigencias.py: hoja Reglas (REGLAS_TABLA, USOS_T*, EXC_*, "
    "GRUPOS_COMUNAS, AGRUPACIONES_TABLAS)",
    "src/hoja_usos.py: TIPOLOGIAS, TABLAS_17C, POR_5C, TABLA_FIJA, ETIQUETAS, NOTAS",
    "src/Liquidacion_tablas.py: asignacion de TABLA_ORIGEN y TABLAS_17C",
    "src/comparacion_vigencia.py: usos_por_modelo y familias",
]


def revisar_reglas(t: pd.DataFrame, mapa: dict) -> list:
    """Avisa si el consolidado o la liquidacion traen algo que las reglas
    escritas no conocen: una tabla nueva, o un uso cuya tabla no aparece en
    el consolidado. No cambia nada; solo imprime la alerta."""
    conocidas = set(TIPOLOGIAS) | set(TABLAS_17C)
    avisos = []
    nuevas = sorted({familia(tabla) for tabla, _ in mapa}
                    - conocidas)
    for fam in nuevas:
        avisos.append(f"El consolidado trae la tabla {fam}, que las reglas "
                      f"no conocen.")
    sin_etiqueta = sorted({familia(tabla) for tabla, _ in mapa}
                          - set(ETIQUETAS) - set(nuevas))
    for fam in sin_etiqueta:
        avisos.append(f"La tabla {fam} no tiene nombre de bloque en ETIQUETAS.")
    for _, f in t.iterrows():
        base = f["TABLA DE VALOR"].split(" · ")[0]
        if base in conocidas and not any(c for _, c in f["_ENLACES"]):
            avisos.append(f"{f['USO_LADM']} ({f['GRUPO']}) se liquida con {base} "
                          f"pero esa tabla no está en el consolidado.")
    if avisos:
        print("\n   " + "!" * 70)
        print("   REVISAR LAS REGLAS: hay tablas o usos que no coinciden")
        for a in avisos:
            print(f"     - {a}")
        print("   Lugares donde se escriben las reglas:")
        for lugar in LUGARES_REGLAS:
            print(f"     * {lugar}")
        print("   " + "!" * 70 + "\n")
    return avisos


HOJA_REVISAR = "REVISAR_REGLAS"
TODAS_COND = set(range(10))
# Tablas que valen para las 17 comunas aunque la hoja de enero escriba solo
# "1,3,9,10,11,12 y 22": las institucionales y las de bloque unico.
FAMILIAS_17C = set(TABLAS_17C) | {"T7_INSTITUCIONAL_SER", "T8_INSTITUCIONAL_IG",
                                   "T12_PARQUEADEROS"}


def _cond_enero(texto):
    """'9', '8 y 9', 'Diferente de 8 y 9', 'Condicion del 1 al 9', 'Todas',
    'Resto' -> conjunto de condiciones (o 'RESTO', o None si viene vacia)."""
    t = str(texto).strip().lower()
    if t in ("", "nan", "none"):
        return None
    if "todas" in t:
        return set(TODAS_COND)
    if "resto" in t:
        return "RESTO"
    m = re.search(r"(\d+)\s*al?\s*(\d+)", t)
    nums = (set(range(int(m.group(1)), int(m.group(2)) + 1)) if m
            else {int(x) for x in re.findall(r"\d+", t)})
    return TODAS_COND - nums if "diferente" in t else nums


def reglas_enero() -> list:
    """La asignacion de la hoja Uso del libro de homologacion: por uso, la
    tabla, las condiciones y las comunas donde aplica."""
    libros = [f for f in sorted(CARPETA_HOMOLOGACION.glob(
        "HOMOLOGACION_Y_TABLAS_DE_VALOR_USOS_LADM_*.xlsx"))
        if not f.name.startswith("~$")]
    if not libros:
        return []
    u = pd.read_excel(libros[-1], sheet_name="Uso", header=1, dtype=str)
    u = u.iloc[:, :11]
    u.columns = ["CAT", "ANT", "NOM", "COD", "USO", "TIPO", "DESC", "COND",
                 "TABLA", "CONDAP", "COMUNA"]
    u["USO"] = u["USO"].ffill().str.strip()
    for c in ["TABLA", "COND", "CONDAP", "COMUNA"]:     # celdas combinadas
        u[c] = u.groupby("USO")[c].ffill()
    u = u[u["TABLA"].notna()].drop_duplicates(["USO", "TABLA", "CONDAP",
                                               "COMUNA"])
    reglas = []
    for _, r in u.iterrows():
        tabla = str(r["TABLA"]).strip()
        fam = "MODELO" if tabla.upper().startswith("MODELO") else familia(tabla)
        cond = _cond_enero(r["CONDAP"])
        texto_cond = r["CONDAP"]
        if cond is None:
            cond, texto_cond = _cond_enero(r["COND"]), r["COND"]
        com = ({x.zfill(2) for x in re.findall(r"\d+", str(r["COMUNA"]))}
               if str(r["COMUNA"]) != "nan" else None)
        com_enero = set(com) if com else set()
        if fam in FAMILIAS_17C:
            com = set(COMUNAS_17)
        condap = str(r["CONDAP"]).strip()
        reglas.append({"uso": r["USO"], "fam": fam, "com_enero": com_enero,
                       "texto_condap": "" if condap.lower() == "nan" else condap,
                       "texto_general": " ".join(str(r["COND"]).split()),
                       "cond": cond if cond is not None else set(TODAS_COND),
                       "com": com, "texto_cond": str(texto_cond).strip(),
                       "texto_com": "17 comunas" if fam in FAMILIAS_17C
                       else str(r["COMUNA"]).strip()})
    for r in reglas:                     # 'Resto' = lo que no cubren las demas
        if r["cond"] == "RESTO":
            otras = [o["cond"] for o in reglas
                     if o is not r and o["uso"] == r["uso"]
                     and o["cond"] != "RESTO"
                     and (o["com"] is None or r["com"] is None
                          or o["com"] & r["com"])]
            r["cond"] = TODAS_COND - set().union(*otras) if otras else set(TODAS_COND)
    return reglas


def comparar_con_enero(liq: pd.DataFrame,
                       mapa: dict | None = None) -> pd.DataFrame:
    """Donde la hoja Uso de enero define una regla, compara la tabla que dice
    contra la que uso la liquidacion. Sin las 5 comunas extra ni rurales."""
    reglas = reglas_enero()
    if not reglas:
        return pd.DataFrame()
    d = liq[["USO_LADM", "CONDICION", "COMUNA", "TABLA_ORIGEN"]].copy()
    d["COMUNA"] = d["COMUNA"].astype(str).str.strip().str.zfill(2)
    d = d[d["COMUNA"].isin(COMUNAS_17)]
    d["USO_LADM"] = d["USO_LADM"].astype(str).str.strip().replace(ALIAS_USOS)
    d["C"] = pd.to_numeric(d["CONDICION"], errors="coerce").fillna(-1).astype(int)
    d["FAM"] = d["TABLA_ORIGEN"].astype(str).map(familia)
    g = d.groupby(["USO_LADM", "COMUNA", "C", "FAM"]).size().reset_index(name="N")

    usos_enero = {r["uso"] for r in reglas}
    filas = []
    for _, o in g.iterrows():
        del_uso = [r for r in reglas if r["uso"] == o["USO_LADM"]]
        if not del_uso:
            filas.append((o["USO_LADM"], "Sin regla en la hoja de enero", "", "",
                          o["FAM"], o["COMUNA"], o["C"], o["N"]))
            continue
        en_comuna = [r for r in del_uso
                     if r["com"] is None or o["COMUNA"] in r["com"]]
        if not en_comuna:
            continue                     # enero no habla de esta comuna
        aplica = [r for r in en_comuna if o["C"] in r["cond"]]
        if not aplica:
            filas.append((o["USO_LADM"], "Condición no prevista en enero",
                          "; ".join(sorted({r["texto_cond"] for r in en_comuna})),
                          "; ".join(sorted({r["fam"] for r in en_comuna})),
                          o["FAM"], o["COMUNA"], o["C"], o["N"]))
        elif ("MODELO" in {r["fam"] for r in aplica}
              and o["FAM"] != "MODELO"):
            # Va por modelo, pero el modelo esta apagado.
            en_consolidado = o["FAM"] in {familia(tb) for tb, _ in (mapa or {})}
            filas.append((o["USO_LADM"],
                          "Va por modelo y la liquidación le pone tabla"
                          if en_consolidado else
                          "Va por modelo y la liquidación le pone una tabla "
                          "que no está en el consolidado",
                          "; ".join(sorted({r["texto_cond"] for r in aplica})),
                          "MODELO", o["FAM"], o["COMUNA"], o["C"], o["N"]))
        elif o["FAM"] not in {r["fam"] for r in aplica}:
            filas.append((o["USO_LADM"], "Tabla distinta",
                          "; ".join(sorted({r["texto_cond"] for r in aplica})),
                          "; ".join(sorted({r["fam"] for r in aplica})),
                          o["FAM"], o["COMUNA"], o["C"], o["N"]))
    if not filas:
        return pd.DataFrame()
    t = pd.DataFrame(filas, columns=["USO_LADM", "DIFERENCIA",
                                     "CONDICIÓN EN ENERO", "TABLA EN ENERO",
                                     "TABLA EN LA LIQUIDACIÓN", "COMUNA",
                                     "CONDICIÓN", "CONSTRUCCIONES"])
    return (t.groupby(["USO_LADM", "DIFERENCIA", "CONDICIÓN EN ENERO",
                       "TABLA EN ENERO", "TABLA EN LA LIQUIDACIÓN"])
             .agg(COMUNAS=("COMUNA", lambda s: ", ".join(sorted(set(s)))),
                  CONDICIONES=("CONDICIÓN",
                               lambda s: ", ".join(map(str, sorted(set(s))))),
                  CONSTRUCCIONES=("CONSTRUCCIONES", "sum"))
             .reset_index()
             .sort_values("CONSTRUCCIONES", ascending=False))


def hoja_revisar(wb, dif: pd.DataFrame) -> None:
    """Las diferencias con la hoja Uso de enero, en una hoja al final."""
    if HOJA_REVISAR in wb.sheetnames:
        del wb[HOJA_REVISAR]
    if dif.empty:
        return
    ws = wb.create_sheet(HOJA_REVISAR)
    titulo = ws.cell(1, 1, "REGLAS QUE NO COINCIDEN: hoja Uso de enero frente a "
                           "lo que hace la liquidación (sin comunas extra ni "
                           "rurales; institucionales y tablas 17C en las 17 "
                           "comunas)")
    ws.merge_cells(start_row=1, start_column=1, end_row=1,
                   end_column=len(dif.columns))
    titulo.font = Font(bold=True, color="FFFFFF", size=12)
    titulo.fill = PatternFill("solid", fgColor="9C3B0B")
    titulo.alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[1].height = 36
    borde = Border(*(Side(style="thin", color="E3C4B0"),) * 4)
    for j, n in enumerate(dif.columns, start=1):
        c = ws.cell(2, j, n)
        c.fill = PatternFill("solid", fgColor="FCE4D6")
        c.font = Font(bold=True, color="0F1F33")
        c.border = borde
        c.alignment = Alignment(horizontal="center", vertical="center",
                                wrap_text=True)
    for i, fila in enumerate(dif.itertuples(index=False), start=3):
        for j, v in enumerate(fila, start=1):
            c = ws.cell(i, j, v)
            c.border = borde
            c.alignment = Alignment(vertical="top", wrap_text=True)
    for col, w in zip("ABCDEFGH", [34, 28, 24, 22, 24, 30, 14, 16]):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A3"
    volver = ws.cell(1, len(dif.columns) + 2, "← Volver a USOS")
    volver.hyperlink = Hyperlink(ref=volver.coordinate, location=f"'{HOJA}'!A1")
    volver.font = Font(bold=True, color="0563C1", underline="single")


HOJA_CONDICION = "CONDICIÓN Y MODELO"


def hoja_condicion_modelo(wb) -> None:
    """Las reglas de condicion juridica y modelo, escritas en una hoja: que va
    por modelo, que va por tabla y que va a preinconsistencias."""
    try:
        from pre_inconsistencias import USOS_PAR_DE_MODELO
    except Exception:                                    # pragma: no cover
        USOS_PAR_DE_MODELO = set()
    if HOJA_CONDICION in wb.sheetnames:
        del wb[HOJA_CONDICION]
    ws = wb.create_sheet(HOJA_CONDICION, 1)
    color = "1D3557"
    borde = Border(*(Side(style="thin", color="B7C9D6"),) * 4)

    def banda(fila, txt):
        c = ws.cell(fila, 1, txt)
        ws.merge_cells(start_row=fila, start_column=1, end_row=fila, end_column=4)
        c.fill = PatternFill("solid", fgColor=color)
        c.font = Font(bold=True, color="FFFFFF", size=12)
        c.alignment = Alignment(vertical="center", indent=1)
        ws.row_dimensions[fila].height = 24

    def tabla(fila, encabezados, filas, tonos=None):
        for j, n in enumerate(encabezados, start=1):
            c = ws.cell(fila, j, n)
            c.fill = PatternFill("solid", fgColor="83CAEB")
            c.font = Font(bold=True, color="0F1F33")
            c.border = borde
            c.alignment = Alignment(horizontal="center", vertical="center",
                                    wrap_text=True)
        for i, valores in enumerate(filas, start=fila + 1):
            for j, v in enumerate(valores, start=1):
                c = ws.cell(i, j, v)
                c.border = borde
                c.alignment = Alignment(vertical="top", wrap_text=True)
                if tonos and tonos[i - fila - 1]:
                    c.fill = PatternFill("solid", fgColor=tonos[i - fila - 1])
            ws.row_dimensions[i].height = 45
        return fila + len(filas) + 2

    ws.cell(1, 1, "CONDICIÓN JURÍDICA Y MODELO · cómo se decide si una "
                  "construcción va por modelo, por tabla o no se liquida"
            ).font = Font(bold=True, size=13, color="0F1F33")
    ws.merge_cells("A1:D1")
    volver = ws.cell(1, 6, "← Volver a USOS")
    volver.hyperlink = Hyperlink(ref=volver.coordinate, location=f"'{HOJA}'!A1")
    volver.font = Font(bold=True, color="0563C1", underline="single")
    ws.cell(2, 1, "La condición jurídica es la posición 22 del número predial: "
                  "8 y 9 son propiedad horizontal (PH).").font = Font(
        italic=True, color="4A5565")
    ws.merge_cells("A2:D2")

    f = 4
    banda(f, "1. Usos que van por modelo")
    f = tabla(f + 1, ["USO (CÓDIGO)", "CONDICIÓN", "SE LIQUIDA CON", "NOTA"], [
        ["Apartamentos_4_y_mas_pisos_en_PH (001)", "9", "MODELO",
         "Mientras el modelo esté apagado queda en $ 0: no se reemplaza con tabla."],
        ["Apartamentos_4_y_mas_pisos_en_PH (001)", "8", "T1_RESIDENCIAL", "Por tabla."],
        ["Comercio_en_PH (028)", "9", "MODELO",
         "Mientras el modelo esté apagado queda en $ 0: no se reemplaza con tabla."],
        ["Comercio_en_PH (028)", "8", "T3_COMERCIAL", "Por tabla."],
        ["Oficinas_Consultorios_en_PH (034)", "9", "MODELO",
         "Mientras el modelo esté apagado queda en $ 0: no se reemplaza con tabla."],
        ["Oficinas_Consultorios_en_PH (034)", "8", "T7_INSTITUCIONAL_SER",
         "Por tabla (la T7 aún no viene en el consolidado)."],
    ], ["E9EEF5", None, "E9EEF5", None, "E9EEF5", None])
    ws.cell(f - 1, 1, "Aplica en todas las comunas. En la comparación de "
                      "vigencias (la app) solo entra la condición 8.").font = Font(
        italic=True, color="4A5565")
    ws.merge_cells(start_row=f - 1, start_column=1, end_row=f - 1, end_column=4)

    f += 1
    banda(f, "2. Uso y condición que no coinciden (preinconsistencias)")
    par = ", ".join(sorted(USOS_PAR_DE_MODELO)) or "—"
    f = tabla(f + 1, ["CASO", "COMUNAS", "QUÉ PASA", "MOTIVO EN EL REPORTE"], [
        ["Uso de PH (nombre termina en _PH) en un predio que no está en "
         "condición 8 ni 9", "Todas", "Se reporta y el predio NO se liquida",
         "USO _PH sin CONDICION 8 o 9"],
        ["Uso de no PH en condición 9", "10 comunas", "Se reporta y el predio NO "
         "se liquida (la condición 9 lo mandaría a la tabla de PH)",
         "USO sin _PH con CONDICION 9"],
        [f"Uso de no PH cuyo par en PH va por modelo ({par}) en condición 9",
         "7 comunas", "Se reporta y el predio NO se liquida: en PH debería ir "
         "por modelo", "USO sin _PH con CONDICION 9 (su par en PH va por modelo)"],
        ["Cualquier otro uso de no PH en condición 9 (depósitos, garajes, "
         "vivienda hasta 3 pisos, pensiones, institucionales…)", "7 comunas",
         "Se reporta pero SÍ se liquida: en esas comunas la tabla no cambia con "
         "la condición", "USO sin _PH con CONDICION 9 (7 comunas, no afecta el valor)"],
    ], ["FCE4D6", "FCE4D6", "FCE4D6", None])
    ws.cell(f - 1, 1, "El detalle de cada caso está en "
                      "output/Pre_inconsistencias/reporte_pre_inconsistencias_"
                      "<fecha>.xlsx, hoja USO_INCONSISTENTE.").font = Font(
        italic=True, color="4A5565")
    ws.merge_cells(start_row=f - 1, start_column=1, end_row=f - 1, end_column=4)

    for col, w in zip("ABCDEF", [46, 16, 46, 52, 4, 18]):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A3"


HOJA_USO_ENERO = "Uso"
GRUPOS_VIEJOS_5 = {"1", "9", "10", "11", "12"}       # el '5C' de enero
GRUPOS_VIEJOS_2 = {"3", "22"}                         # el '2C' de enero
FAMILIAS_TODAS = (set(TABLAS_17C) | {"T7_INSTITUCIONAL_SER", "T8_INSTITUCIONAL_IG",
                                     "T12_PARQUEADEROS", "MODELO"})


def _grupos_nuevos(fam: str, viejas: set) -> list:
    """Los grupos de comunas de hoy para una fila de enero."""
    if fam in FAMILIAS_TODAS:
        return ["17C"]
    if fam in TIPOLOGIAS or fam == "T10_ANEXOS":
        grupos = []
        if viejas & GRUPOS_VIEJOS_5:
            grupos += ["5C", "5C_N"] if fam in POR_5C else ["10C"]
        if viejas & GRUPOS_VIEJOS_2:
            grupos.append("7C")
        return grupos
    return []


def hoja_usos_enero(wb, bloques: dict, bloque_t12: dict | None) -> tuple:
    """La hoja Uso del libro de homologacion de enero, TAL CUAL (filas, textos,
    celdas combinadas y formato). Solo cambian la COMUNA, con los grupos de
    hoy, y el HIPERVINCULO, al bloque de la tabla en el consolidado."""
    from copy import copy
    libros = [f for f in sorted(CARPETA_HOMOLOGACION.glob(
        "HOMOLOGACION_Y_TABLAS_DE_VALOR_USOS_LADM_*.xlsx"))
        if not f.name.startswith("~$")]
    if not libros:
        raise RuntimeError("No esta el libro de homologacion en input/tablas.")
    en = load_workbook(libros[-1])[HOJA_USO_ENERO]

    if HOJA in wb.sheetnames:
        del wb[HOJA]
    ws = wb.create_sheet(HOJA, 0)

    # Hasta la ultima fila con datos en A:M; N en adelante enlazaba a hojas
    # del libro de enero que aqui no existen.
    ultima = max(c.row for fila in en.iter_rows(max_col=13) for c in fila
                 if c.value not in (None, ""))
    for fila in en.iter_rows(min_row=1, max_row=ultima, max_col=13):
        for c in fila:
            n = ws.cell(c.row, c.column, c.value)
            if c.has_style:
                n.font, n.fill, n.border = copy(c.font), copy(c.fill), copy(c.border)
                n.alignment, n.number_format = copy(c.alignment), c.number_format
    for m in en.merged_cells.ranges:
        if m.max_col <= 13 and m.max_row <= ultima:
            ws.merge_cells(str(m))
    for letra, dim in en.column_dimensions.items():
        if dim.width and letra in "ABCDEFGHIJKLM":
            ws.column_dimensions[letra].width = dim.width
    for r in range(1, ultima + 1):
        if en.row_dimensions[r].height:
            ws.row_dimensions[r].height = en.row_dimensions[r].height

    # Valor efectivo de cada celda: en las combinadas, el de su esquina.
    esquina = {}
    for m in en.merged_cells.ranges:
        for r in range(m.min_row, m.max_row + 1):
            for col in range(m.min_col, m.max_col + 1):
                esquina[(r, col)] = (m.min_row, m.min_col)

    def valor(r, col):
        rr, cc = esquina.get((r, col), (r, col))
        v = en.cell(rr, cc).value
        return "" if v is None else str(v).strip()

    def es_esquina(r, col):
        return esquina.get((r, col), (r, col)) == (r, col)

    K, L = 11, 12
    n_links = 0
    for r in range(3, ultima + 1):
        tabla, viejas = valor(r, 9), valor(r, K)
        if not tabla or not es_esquina(r, K):
            continue
        fam = "MODELO" if tabla.upper().startswith("MODELO") else familia(tabla)
        grupos = _grupos_nuevos(fam, set(re.findall(r"\d+", viejas)))
        if grupos:
            ws.cell(r, K).value = " · ".join(f"{g}: {comunas_de(g)}"
                                             for g in grupos)
        # El bloque del consolidado (o la hoja T12) al que lleva la fila.
        bloque = None
        if fam == "T12_PARQUEADEROS":
            bloque = bloque_t12
        elif fam in TABLAS_17C:
            bloque = bloques.get(f"{fam}_17C")
        elif fam in TIPOLOGIAS and grupos:
            g = grupos[0]
            if fam == "T1_RESIDENCIAL" and g == "10C":
                cond = valor(r, 10).lower()
                g = ("10C_COND_9" if "9" in cond and "diferente" not in cond
                     and "resto" not in cond else "10C_COND_0")
            bloque = bloques.get(f"{fam}_{g}")
        if not es_esquina(r, L):
            continue
        c = ws.cell(r, L)
        c.hyperlink = None
        if bloque:
            c.value = f"{bloque['etiqueta']} →"
            c.hyperlink = Hyperlink(ref=c.coordinate,
                                    location=f"'{bloque.get('hoja', HOJA_CONV)}'!"
                                             f"{bloque['celda']}")
            c.fill = PatternFill("solid", fgColor=bloque["color"])
            c.font = Font(bold=True, underline="single", color=bloque["texto"])
            n_links += 1
        else:
            c.value = None
    ws.cell(1, 13, "Reglas de condición y modelo →").hyperlink = Hyperlink(
        ref="M1", location=f"'{HOJA_CONDICION}'!A1")
    ws.cell(1, 13).font = Font(bold=True, color="0563C1", underline="single")
    ws.freeze_panes = "F3"
    return ultima - 2, n_links


HOJA_T12 = "T12_PARQUEADEROS"
HOJA_T12_ENERO = "T12_PARQUEADEROS_7C"
COLOR_T12 = "566573"


def hoja_t12(wb) -> dict | None:
    """La hoja T12 del libro de homologacion de enero, con el MISMO contenido
    (textos, valores y orden) y mejor presentacion. Devuelve el bloque para
    enlazarlo desde USOS, o None si el libro no esta."""
    from consolidar_tablas import _mezclar, _texto_contraste
    libros = [f for f in sorted(CARPETA_HOMOLOGACION.glob(
        "HOMOLOGACION_Y_TABLAS_DE_VALOR_USOS_LADM_*.xlsx"))
        if not f.name.startswith("~$")]
    if not libros:
        print("   (sin libro de homologacion: la hoja T12 no se arma)")
        return None
    origen = load_workbook(libros[-1], data_only=True)
    if HOJA_T12_ENERO not in origen.sheetnames:
        print(f"   (el libro de homologacion no trae {HOJA_T12_ENERO})")
        return None
    en = origen[HOJA_T12_ENERO]

    def val(coord):
        v = en[coord].value
        return v.strip() if isinstance(v, str) else v

    if HOJA_T12 in wb.sheetnames:
        del wb[HOJA_T12]
    ws = wb.create_sheet(HOJA_T12, wb.sheetnames.index(HOJA_CONV) + 1)
    color = COLOR_T12
    texto = _texto_contraste(color)
    borde = Border(*(Side(style="thin", color="B7C9D6"),) * 4)
    pesos = '"$" #,##0'
    tonos = {"Residencial": _mezclar(color, 0.88), "Comercio": _mezclar(color, 0.78),
             "Deposito": _mezclar(color, 0.68)}
    ancho = 7                                            # columnas A..G

    def banda(fila, txt, hasta=ancho, size=12):
        c = ws.cell(fila, 1, txt)
        ws.merge_cells(start_row=fila, start_column=1, end_row=fila,
                       end_column=hasta)
        c.fill = PatternFill("solid", fgColor=color)
        c.font = Font(bold=True, color=texto, size=size)
        c.alignment = Alignment(vertical="center", indent=1)
        ws.row_dimensions[fila].height = 24

    def encabezado(fila, nombres):
        for j, n in enumerate(nombres, start=1):
            c = ws.cell(fila, j, n)
            c.fill = PatternFill("solid", fgColor=_mezclar(color, 0.55))
            c.font = Font(bold=True, color="0F1F33")
            c.border = borde
            c.alignment = Alignment(horizontal="center", vertical="center",
                                    wrap_text=True)
        ws.row_dimensions[fila].height = 32

    def celda(fila, col, v, fondo=None, moneda=False, negrilla=False,
              centro=False):
        c = ws.cell(fila, col, v)
        c.border = borde
        c.alignment = Alignment(vertical="center", wrap_text=True,
                                horizontal="center" if centro else None)
        if fondo:
            c.fill = PatternFill("solid", fgColor=fondo)
        if moneda and isinstance(v, (int, float)):
            c.number_format = pesos
            c.alignment = Alignment(horizontal="right", vertical="center")
        if negrilla:
            c.font = Font(bold=True)
        return c

    def texto_libre(fila, txt, italica=False):
        c = ws.cell(fila, 1, txt)
        ws.merge_cells(start_row=fila, start_column=1, end_row=fila,
                       end_column=ancho)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        c.font = Font(italic=italica, color="4A5565" if italica else "1B2430")
        ws.row_dimensions[fila].height = 30

    # --- Titulo ----------------------------------------------------------------
    banda(1, "T12 · PARQUEADEROS Y DEPÓSITOS", size=14)
    volver = ws.cell(1, ancho + 2, "← Volver a USOS")
    volver.hyperlink = Hyperlink(ref=volver.coordinate, location=f"'{HOJA}'!A1")
    volver.font = Font(bold=True, color="0563C1", underline="single")

    # --- Paso 1 ----------------------------------------------------------------
    banda(3, val("B1"))
    for k, r in enumerate((3, 4), start=4):
        celda(k, 1, val(f"B{r}"), _mezclar(color, 0.88), negrilla=True)
        c = celda(k, 2, val(f"C{r}"))
        ws.merge_cells(start_row=k, start_column=2, end_row=k, end_column=ancho)

    # --- Paso 2 ----------------------------------------------------------------
    banda(7, val("B6"))
    texto_libre(8, val("B7"), italica=True)
    texto_libre(9, val("B8"), italica=True)
    encabezado(10, [val(f"{c}9") for c in "BCDEFGH"])
    fila = 11
    for r in range(10, 36):
        datos = [val(f"{c}{r}") for c in "BCDEFGH"]
        if not any(x not in (None, "") for x in datos):
            continue
        grupo = str(datos[0] or "")
        fondo = next((t for g, t in tonos.items() if grupo.startswith(g)), None)
        for j, v in enumerate(datos, start=1):
            celda(fila, j, v, fondo, moneda=(j == 6), negrilla=(j == 1),
                  centro=j in (3, 4, 5))
        ws.row_dimensions[fila].height = 30
        fila += 1
    fin_tabla = fila

    # --- Calculo del factor (tabla de enero, filas 43 a 49) --------------------
    f = fin_tabla + 1
    banda(f, "Cálculo del factor")
    encabezado(f + 1, [val(f"{c}43") for c in "BCDE"])
    inicio = f + 2
    for k, r in enumerate(range(44, 50)):
        for j, c in enumerate("BCDE", start=1):
            celda(inicio + k, j, val(f"{c}{r}"), centro=j in (3,),
                  negrilla=(j == 1))
        ws.row_dimensions[inicio + k].height = 30
    # Mismas celdas combinadas que en enero: 007/008 en tres filas, 036/037 en dos.
    ws.merge_cells(start_row=inicio, start_column=1, end_row=inicio + 2,
                   end_column=1)
    ws.merge_cells(start_row=inicio + 3, start_column=1, end_row=inicio + 4,
                   end_column=1)
    for r in range(inicio, inicio + 6):
        ws.cell(r, 1).alignment = Alignment(wrap_text=True, vertical="center")
    # La fila de 006/027 lleva dos lineas largas en la primera columna.
    ws.row_dimensions[inicio + 5].height = 72.6

    # --- Paso 4 ----------------------------------------------------------------
    f = inicio + 7
    banda(f, val("B38"))
    texto_libre(f + 1, f"• {val('B39')}")
    texto_libre(f + 2, f"• {val('B40')}")

    for col, w in zip("ABCDEFGHI", [30, 29.22, 18, 12, 24, 18, 44, 4, 18]):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"
    return {"etiqueta": "PARQUEADEROS Y DEPÓSITOS", "celda": "A1",
            "hoja": HOJA_T12, "color": color, "texto": texto,
            "claro": _mezclar(color, 0.88)}


if __name__ == "__main__":
    agregar_hoja_usos()
